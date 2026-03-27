# app/parser_excel.py
import os
from openpyxl import load_workbook
from typing import List, Dict, Any, Optional


def parse_excel_files(file_paths: List[str]) -> List[Dict[str, Any]]:
    """
    Парсит список Excel-файлов и возвращает список словарей с информацией о формулах.

    :param file_paths: Список путей к Excel-файлам
    :return: Список словарей с данными о формулах
    """
    all_formulas = []

    for file_path in file_paths:
        try:
            formulas = extract_formulas_from_file(file_path)
            all_formulas.extend(formulas)
        except Exception as e:
            print(f"[ERROR] Ошибка при обработке файла {file_path}: {str(e)}")
    
    return all_formulas


def extract_formulas_from_file(file_path: str) -> List[Dict[str, Any]]:
    """
    Извлекает формулы и значения из одного Excel-файла.

    :param file_path: Путь к Excel-файлу
    :return: Список словарей с информацией о формулах
    """
    results = []
    filename = os.path.basename(file_path)
    
    # Открываем книгу дважды: один раз для формул, второй - для значений
    wb_formulas = load_workbook(file_path, data_only=False)
    wb_values = load_workbook(file_path, data_only=True)
    
    for sheet_name in wb_formulas.sheetnames:
        ws_formulas = wb_formulas[sheet_name]
        ws_values = wb_values[sheet_name]
        
        # Перебираем только ячейки с формулами
        for row in ws_formulas.iter_rows():
            for cell in row:
                if cell.value is not None and isinstance(cell.value, str) and cell.value.startswith('='):
                    coord = cell.coordinate
                    formula = cell.value
                    
                    # Получаем значение из второй версии книги
                    value_cell = ws_values[cell.coordinate]
                    value = value_cell.value
                    
                    # Проверяем на внешние ссылки
                    is_external = has_external_link(formula)
                    
                    results.append({
                        'file': filename,
                        'sheet': sheet_name,
                        'cell': coord,
                        'formula': formula,
                        'value': value,
                        'is_external': is_external
                    })
    
    return results


def has_external_link(formula: str) -> bool:
    """
    Проверяет, содержит ли формула внешние ссылки.

    :param formula: Формула для проверки
    :return: True, если есть внешняя ссылка
    """
    return (
        ('[' in formula and ']' in formula) or
        '.xls' in formula or
        '.xlsx' in formula
    )


def save_formulas_to_txt(formulas: List[Dict[str, Any]], output_path: str) -> None:
    """
    Сохраняет список формул в текстовый файл.

    :param formulas: Список словарей с формулами
    :param output_path: Путь к выходному файлу
    """
    with open(output_path, 'w', encoding='utf-8') as f:
        for item in formulas:
            line = (
                f"[Book]: {item['file']} "
                f"[Sheet]: {item['sheet']} "
                f"[Cell]: {item['cell']} "
                f"[Value]: {item['value']} "
                f"[Formula]: {item['formula']}"
            )
            if item['is_external']:
                line += " [EXTERNAL LINK]"
            f.write(line + '\n')
