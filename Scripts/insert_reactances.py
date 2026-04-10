import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from app.models import Reactance, Base


DATABASE_URL = "mssql+pyodbc:///?odbc_connect=DRIVER={ODBC Driver 18 for SQL Server};SERVER=MSI-ALEXNAB\\SQLEXPRESS;DATABASE=RZA_Calculator;Trusted_Connection=yes;TrustServerCertificate=yes"
ENGINE = create_engine(DATABASE_URL)


def _find_expert_workbook_path() -> Path:
    p = Path("data/input")
    files = [f for f in p.glob("*.xlsx") if not f.name.startswith("~$")]
    if not files:
        raise RuntimeError("No .xlsx files found in data/input")
    return min(files, key=lambda x: x.stat().st_size)


def _find_reactances_sheet(wb):
    if "Реактансы" in wb.sheetnames:
        return wb["Реактансы"]
    for name in wb.sheetnames:
        ws = wb[name]
        c1 = ws.cell(1, 3).value
        d1 = ws.cell(1, 4).value
        if isinstance(c1, str) and isinstance(d1, str) and "MAX" in c1.upper() and "MIN" in d1.upper():
            return ws
    raise RuntimeError("Reactances sheet not found (Реактансы / MAX/MIN)")


def upsert_reactances_from_expert() -> int:
    Base.metadata.create_all(ENGINE)

    expert_path = _find_expert_workbook_path()
    wb = load_workbook(expert_path, data_only=True)
    ws = _find_reactances_sheet(wb)

    count = 0
    with Session(ENGINE) as session:
        session.query(Reactance).delete()
        session.commit()

        for r in range(2, 2000):
            code = ws.cell(r, 1).value
            name = ws.cell(r, 2).value
            zmax = ws.cell(r, 3).value
            zmin = ws.cell(r, 4).value
            z_a_max = ws.cell(r, 10).value
            z_a_min = ws.cell(r, 11).value

            if code is None and name is None:
                break
            if not isinstance(code, (int, float)) or not name:
                continue

            has_n = isinstance(zmax, (int, float)) and isinstance(zmin, (int, float))
            has_a = (
                isinstance(z_a_max, (int, float))
                and isinstance(z_a_min, (int, float))
                and float(z_a_max) > 0
                and float(z_a_min) > 0
            )
            if not has_n and not has_a:
                continue

            session.add(
                Reactance(
                    reactance_code=int(code),
                    name=str(name),
                    z_max_ohm=float(zmax) if has_n else None,
                    z_min_ohm=float(zmin) if has_n else None,
                    z_a_max_ohm=float(z_a_max) if has_a else None,
                    z_a_min_ohm=float(z_a_min) if has_a else None,
                    is_active=True,
                )
            )
            count += 1

        session.commit()

    return count


def main() -> None:
    count = upsert_reactances_from_expert()
    print(f"Inserted reactances: {count}")


if __name__ == "__main__":
    main()
