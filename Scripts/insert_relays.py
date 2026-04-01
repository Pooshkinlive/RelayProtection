import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from app.models import RelayType, Base


DATABASE_URL = "mssql+pyodbc:///?odbc_connect=DRIVER={ODBC Driver 18 for SQL Server};SERVER=MSI-ALEXNAB\\SQLEXPRESS;DATABASE=RZA_Calculator;Trusted_Connection=yes;TrustServerCertificate=yes"
ENGINE = create_engine(DATABASE_URL)


def upsert_relays_from_expert_xlsx(path: str) -> int:
    Base.metadata.create_all(ENGINE)

    wb = load_workbook(path, data_only=True)

    # In this workbook, relay sheet is named "Реле".
    # IMPORTANT: do NOT auto-pick a sheet by generic A2/B2/C2/D2 shape, because
    # "Провода"/"Кабели" match that shape too and pollute relay_types.
    relay_ws = None

    # 1) Prefer exact name.
    if "Реле" in wb.sheetnames:
        relay_ws = wb["Реле"]

    # 2) Fallback: detect by headers in row 1 containing "МТО" and "МТЗ"
    # (these are characteristic for the relay sheet in ЭКСПЕРТ.xlsx).
    if relay_ws is None:
        for ws in wb.worksheets:
            h1 = ws.cell(1, 3).value
            h2 = ws.cell(1, 4).value
            if isinstance(h1, str) and isinstance(h2, str):
                if "МТО" in h1 and "МТЗ" in h2:
                    relay_ws = ws
                    break

    # 3) Fallback: scan first few rows for relay name containing "РТ" or "PT"
    if relay_ws is None:
        for ws in wb.worksheets:
            for r in range(2, 15):
                b = ws.cell(r, 2).value
                if isinstance(b, str) and (("РТ" in b) or ("PT" in b)):
                    relay_ws = ws
                    break
            if relay_ws is not None:
                break

    if relay_ws is None:
        raise RuntimeError("Relay sheet not found in ЭКСПЕРТ.xlsx")

    count = 0
    with Session(ENGINE) as session:
        # Clear old wrong data (e.g., if it was populated from 'Провода').
        session.query(RelayType).delete()
        session.commit()

        for r in range(2, 200):  # safe upper bound
            relay_code = relay_ws.cell(r, 1).value
            relay_name = relay_ws.cell(r, 2).value
            coef_l20 = relay_ws.cell(r, 3).value
            coef_l19 = relay_ws.cell(r, 4).value
            if relay_code is None and relay_name is None:
                break
            if not isinstance(relay_code, (int, float)) or relay_name is None:
                continue
            if not isinstance(coef_l20, (int, float)) or not isinstance(coef_l19, (int, float)):
                continue

            relay_code_i = int(relay_code)
            session.add(
                RelayType(
                    relay_code=relay_code_i,
                    relay_name=str(relay_name),
                    coef_l20=float(coef_l20),  # L20: "Для КЗ за трансформатором"
                    coef_l19=float(coef_l19),  # L19: "коэф. намагничивания"
                )
            )
            count += 1

        session.commit()

    return count


def main() -> None:
    path = "data/input/ЭКСПЕРТ.xlsx"
    count = upsert_relays_from_expert_xlsx(path)
    print(f"Inserted/updated relays: {count}")


if __name__ == "__main__":
    main()

