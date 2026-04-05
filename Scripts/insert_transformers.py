import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from openpyxl import load_workbook
import re

from app.models import Transformer, Base


DATABASE_URL = "mssql+pyodbc:///?odbc_connect=DRIVER={ODBC Driver 18 for SQL Server};SERVER=MSI-ALEXNAB\\SQLEXPRESS;DATABASE=RZA_Calculator;Trusted_Connection=yes;TrustServerCertificate=yes"
ENGINE = create_engine(DATABASE_URL)


def _find_expert_workbook_path() -> Path:
    p = Path("data/input")
    files = [f for f in p.glob("*.xlsx") if not f.name.startswith("~$")]
    if not files:
        raise RuntimeError("No .xlsx files found in data/input")
    # pick smallest non-lock (in case of multiple uploads)
    return min(files, key=lambda x: x.stat().st_size)


def _find_transformers_sheet(wb):
    if "Трансформаторы" in wb.sheetnames:
        return wb["Трансформаторы"]
    # Fallback: first sheet where A1/B1/C1 contain transformer-like headers
    for name in wb.sheetnames:
        ws = wb[name]
        a1 = ws.cell(1, 1).value
        b1 = ws.cell(1, 2).value
        c1 = ws.cell(1, 3).value
        candidates = [str(a1).lower(), str(b1).lower(), str(c1).lower()]
        if any("транс" in s for s in candidates) and any("ом" in s for s in candidates):
            return ws
    raise RuntimeError("Transformers sheet not found (Трансформаторы)")


def upsert_transformers_from_expert() -> int:
    Base.metadata.create_all(ENGINE)

    expert_path = _find_expert_workbook_path()
    wb = load_workbook(expert_path, data_only=True)
    ws = _find_transformers_sheet(wb)

    count = 0
    with Session(ENGINE) as session:
        session.query(Transformer).delete()
        session.commit()

        # Expected columns:
        #   A: order code (1..11)
        #   B: power (kVA)
        #   C: Z (Ohm)
        def _to_float(v):
            if v is None:
                return None
            if isinstance(v, (int, float)):
                return float(v)
            if isinstance(v, str):
                # Excel sometimes contains text like "100 кВА" or garbled symbols.
                # Keep only the first numeric token.
                m = re.search(r'-?\d+(?:[.,]\d+)?', v.replace(',', '.'))
                if not m:
                    return None
                return float(m.group(0))
            return None

        for r in range(2, 500):
            code = ws.cell(r, 1).value
            power_raw = ws.cell(r, 2).value
            z_raw = ws.cell(r, 3).value
            power = _to_float(power_raw)
            z = _to_float(z_raw)

            if code is None and power is None and z is None:
                break
            if not isinstance(code, (int, float)):
                continue
            if power is None or z is None:
                continue

            session.add(
                Transformer(
                    transformer_code=int(code),
                    power_kva=float(power),
                    z_ohm=float(z),
                )
            )
            count += 1

        session.commit()

    return count


def main() -> None:
    count = upsert_transformers_from_expert()
    print(f"Inserted transformers: {count}")


if __name__ == "__main__":
    main()

